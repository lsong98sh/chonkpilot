#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""文档转换服务自测（无需启动服务，直接调用 server.py 的转换核心）。

做两件事：

1. **生成中文样本并断言抽取质量**：用 python-docx / openpyxl / python-pptx 现场生成
   中文 docx / xlsx / pptx 到 `samples/generated/`，逐个转换，打印抽取文本与 loc，
   断言关键中文词句确实被抽出（含 xlsx 的 sheet 定位、pptx 的 slide 定位）。
2. **真实文档质量基线（可选）**：若 `samples/` 下存在真实 PDF / Office 文件
   （不递归 `generated/`），逐个转换并打印结果摘要，不做断言，供人工评估。

另含少量负例断言：旧版二进制（.doc）、超大输入（too_large）、路径不存在（not_found）。

用法（在 src/mcps/markitdown 内）：
    .\\.venv\\Scripts\\python.exe selftest.py
"""

from __future__ import annotations

import asyncio
import os
import re
import shutil
import sys
import tempfile
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from server import (  # noqa: E402
    ALLOWED_ROOTS_ENV,
    ALLOW_PRIVATE_HOSTS,
    ERROR_NOT_FOUND,
    ERROR_TOO_LARGE,
    ERROR_UNSUPPORTED_FORMAT,
    LOC_KIND_NONE,
    LOC_KIND_PAGE,
    LOC_KIND_SHEET,
    LOC_KIND_SLIDE,
    SUPPORTED_EXTENSIONS,
    allowed_convert_roots,
    check_remote_host,
    convert_path,
    path_within_roots,
    remove_state,
    run_in_convert_pool,
    write_state,
)

try:  # 控制台可能不是 UTF-8，强制 UTF-8 以便打印中文
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # pragma: no cover
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLES_DIR = os.path.join(HERE, "samples")
GENERATED_DIR = os.path.join(SAMPLES_DIR, "generated")

_failures: list[str] = []


def check(condition: bool, label: str) -> None:
    """记录一条断言结果。"""
    if condition:
        print(f"  [OK]   {label}")
    else:
        print(f"  [FAIL] {label}")
        _failures.append(label)


# ─────────────────────────── 中文样本生成 ───────────────────────────


def build_docx(path: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("文档转换自测报告", level=1)
    doc.add_paragraph(
        "这是用于验证中文抽取质量的一段正文，包含关键字：全文索引、向量检索。"
    )
    doc.add_paragraph("第二条：Office 与 PDF 文档都应可被抽取为 Markdown。")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "名称"
    table.rows[0].cells[1].text = "说明"
    row = table.add_row()
    row.cells[0].text = "中文表格单元"
    row.cells[1].text = "markitdown 转换结果"
    doc.save(path)


def build_xlsx(path: str) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "第一张表"
    sheet.append(["项目", "备注"])
    sheet.append(["服务器采购", "含三年维保"])

    second = workbook.create_sheet("数据表")
    second.append(["指标", "本期"])
    second.append(["活跃用户", "4321"])
    workbook.save(path)


def build_pptx(path: str) -> None:
    from pptx import Presentation

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = "季度汇报"
    title_slide.placeholders[1].text = "第三季度进展与下季度计划"

    content_slide = prs.slides.add_slide(prs.slide_layouts[1])
    content_slide.shapes.title.text = "检索能力升级"
    content_slide.placeholders[1].text = "支持 Office 与 PDF 的全文索引"
    prs.save(path)


def build_pdf(path: str) -> None:
    """生成 2 页带文本层的 ASCII PDF（手写最小 PDF，避免额外依赖），覆盖 PDF 页码定位。

    注：手写 PDF 不便内嵌中文字体（需 CID 字体），故只用 ASCII 文本；
    中文抽取质量由 docx / xlsx / pptx 三个样本覆盖。
    """
    streams = [
        b"BT /F1 18 Tf 72 720 Td (Alpha page one chonkpilot) Tj ET",
        b"BT /F1 18 Tf 72 720 Td (Beta page two chonkpilot) Tj ET",
    ]
    page_obj = (
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 7 0 R >> >> /Contents %d 0 R >>"
    )
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 5 0 R] /Count 2 >>",
        page_obj % 4,
        b"<< /Length %d >>\nstream\n" % len(streams[0]) + streams[0] + b"\nendstream",
        page_obj % 6,
        b"<< /Length %d >>\nstream\n" % len(streams[1]) + streams[1] + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % index + body + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_pos,
    )
    with open(path, "wb") as fh:
        fh.write(bytes(out))


# ─────────────────────────── 打印与断言 ───────────────────────────


def show(outcome, text_limit: int = 800) -> None:
    if outcome.ok:
        print(
            f"  ok=True loc_kind={outcome.loc_kind} locs={len(outcome.locs)} "
            f"truncated={outcome.truncated} elapsed={outcome.elapsed_ms}ms"
        )
        if outcome.locs:
            preview = ", ".join(f"{entry['loc']}@{entry['offset']}" for entry in outcome.locs[:8])
            print(f"  locs: {preview}")
        print("  --- text ---")
        for line in outcome.text[:text_limit].splitlines():
            print(f"  | {line}")
        if len(outcome.text) > text_limit:
            print(f"  | ... (total {len(outcome.text)} chars)")
        print("  --- /text ---")


def check_locs_byte_offsets(title: str, outcome) -> None:
    """断言 locs[].offset = **UTF-8 字节偏移**：必须落在字符边界（中文样本下字符偏移会越界）。"""
    if not outcome.ok or not outcome.locs:
        return
    raw = outcome.text.encode("utf-8")
    for entry in outcome.locs:
        off = entry["offset"]
        in_range = 0 <= off <= len(raw)
        check(in_range, f"{title}: loc {entry['loc']} offset 在字节范围内（{off}）")
        if not in_range:
            continue
        try:
            raw[:off].decode("utf-8")
            boundary = True
        except UnicodeDecodeError:
            boundary = False
        check(boundary, f"{title}: loc {entry['loc']} offset 落在 UTF-8 字符边界（字节口径）")


def positive_case(title: str, path: str, keywords: list[str], expect_kind: str) -> None:
    print(f"\n===== {title} =====")
    print(f"  file: {os.path.relpath(path, HERE)}")
    outcome = convert_path(path)
    show(outcome)
    check(outcome.ok, f"{title}: convert ok")
    if not outcome.ok:
        return
    for keyword in keywords:
        check(keyword in outcome.text, f"{title}: contains '{keyword}'")
    check(outcome.loc_kind == expect_kind, f"{title}: loc_kind == {expect_kind}")
    check_locs_byte_offsets(title, outcome)
    if expect_kind == LOC_KIND_PAGE:
        pages = [entry["loc"] for entry in outcome.locs]
        check(pages == ["1", "2"], f"{title}: page locs == ['1','2'] ({pages})")
    if expect_kind == LOC_KIND_SHEET:
        names = [entry["loc"] for entry in outcome.locs]
        check(len(names) >= 2, f"{title}: >=2 sheet locs ({names})")
        # 中文 sheet 名 → 偏移必须是字节口径（第二张表起点 = 其前文本的 UTF-8 字节长度）
        marks = list(re.finditer(r"^##[ \t]+(.+?)[ \t]*$", outcome.text, re.MULTILINE))
        check(len(marks) == len(outcome.locs), f"{title}: sheet 标记数与 locs 数一致")
        if len(marks) == len(outcome.locs) and marks:
            want = [len(outcome.text[:m.start()].encode("utf-8")) for m in marks]
            got = [entry["offset"] for entry in outcome.locs]
            check(got == want, f"{title}: sheet offset == UTF-8 字节偏移（got={got} want={want}）")
    if expect_kind == LOC_KIND_SLIDE:
        check(len(outcome.locs) >= 2, f"{title}: >=2 slide locs")
    if expect_kind == LOC_KIND_NONE:
        check(not outcome.locs, f"{title}: no locs")


def negative_cases() -> None:
    print("\n===== 负例（错误码） =====")

    legacy = os.path.join(GENERATED_DIR, "legacy.doc")
    with open(legacy, "wb") as fh:
        fh.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    outcome = convert_path(legacy)
    print(f"  legacy.doc -> ok={outcome.ok} error_code={outcome.error_code}")
    check(
        (not outcome.ok) and outcome.error_code == ERROR_UNSUPPORTED_FORMAT,
        f"legacy .doc -> {ERROR_UNSUPPORTED_FORMAT}",
    )

    docx = os.path.join(GENERATED_DIR, "chinese_sample.docx")
    outcome = convert_path(docx, max_input_bytes=10)
    print(f"  too_large  -> ok={outcome.ok} error_code={outcome.error_code}")
    check(
        (not outcome.ok) and outcome.error_code == ERROR_TOO_LARGE,
        f"oversize input -> {ERROR_TOO_LARGE}",
    )

    outcome = convert_path(os.path.join(GENERATED_DIR, "no-such-file.pdf"))
    print(f"  missing    -> ok={outcome.ok} error_code={outcome.error_code}")
    check(
        (not outcome.ok) and outcome.error_code == ERROR_NOT_FOUND,
        f"missing path -> {ERROR_NOT_FOUND}",
    )


def lifecycle_cases() -> None:
    """状态文件写入 / 优雅退出删除 / 不误删他人状态文件。"""
    print("\n===== 生命周期（状态文件） =====")
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "state.json")
        write_state(path, {"pid": os.getpid(), "token": "selftest"})
        check(os.path.exists(path), "write_state 写出 state.json")

        remove_state(path)
        check(not os.path.exists(path), "remove_state（本进程 pid）→ 已删除")

        write_state(path, {"pid": os.getpid() + 1, "token": "selftest"})
        remove_state(path)
        check(os.path.exists(path), "remove_state（他人 pid）→ 不删除")


def security_cases() -> None:
    """安全约束单测（SSRF host 校验 + /vfts/convert 路径根约束）——纯函数，不访问网络。"""
    print("\n===== 安全约束（SSRF / 路径根） =====")

    # SSRF：远端抓取前拒绝解析到本机 / 私网 / 链路本地 / 云元数据（169.254.169.254）地址。
    if ALLOW_PRIVATE_HOSTS:
        print("  （MARKITDOWN_ALLOW_PRIVATE_HOSTS 已开启 → 跳过 SSRF 拒绝断言）")
    else:
        for host in ("127.0.0.1", "localhost", "10.0.0.1", "192.168.1.1",
                     "172.16.0.1", "169.254.169.254", "::1"):
            check(bool(check_remote_host(host)), f"SSRF: 拒绝 {host}")

    # 路径根约束：realpath 后再比较（含符号链接），拒绝越界读。
    root = tempfile.mkdtemp(prefix="markitdown-root-")
    try:
        inside = os.path.join(root, "a.docx")
        outside = os.path.join(os.path.dirname(root), "outside.docx")
        check(path_within_roots(inside, [root]), "root: 根内文件放行")
        check(not path_within_roots(outside, [root]), "root: 根外文件拒绝")
        check(not path_within_roots(inside, []), "root: 无允许根 → 拒绝")
        # 前缀相似目录不得误判（<root>X 不是 <root> 的子目录）
        check(not path_within_roots(os.path.join(root + "X", "a.docx"), [root]),
              "root: 前缀相似目录不误判")
        check(allowed_convert_roots(root) == [root], "root: 请求体 root 计入允许根")

        # 环境白名单按 os.pathsep 解析
        saved = os.environ.get(ALLOWED_ROOTS_ENV)
        os.environ[ALLOWED_ROOTS_ENV] = root + os.pathsep + os.path.dirname(root)
        try:
            check(allowed_convert_roots() == [root, os.path.dirname(root)],
                  f"{ALLOWED_ROOTS_ENV}: 按 pathsep 解析多根")
        finally:
            if saved is None:
                os.environ.pop(ALLOWED_ROOTS_ENV, None)
            else:
                os.environ[ALLOWED_ROOTS_ENV] = saved
    finally:
        shutil.rmtree(root, ignore_errors=True)


def timeout_case() -> None:
    """/vfts/convert 超时控制：run_in_convert_pool 到点抛 TimeoutError，且**不等待线程结束**
    （超时分支已尽力 future.cancel()；线程不可强杀，仅回收未出队任务）。"""
    print("\n===== 转换超时（超时取消） =====")

    def slow() -> str:
        time.sleep(0.6)  # 远超下面的 timeout
        return "late"

    async def run() -> float | None:
        t0 = time.perf_counter()
        try:
            await run_in_convert_pool(slow, timeout=0.05)
        except asyncio.TimeoutError:
            return time.perf_counter() - t0
        return None

    elapsed = asyncio.run(run())
    check(elapsed is not None, "超时应抛 asyncio.TimeoutError")
    if elapsed is not None:
        check(elapsed < 0.5, f"超时应即时返回（不等线程结束）：elapsed={elapsed:.3f}s")

    # 未超时：正常返回结果。
    async def fast_run() -> str:
        return await run_in_convert_pool(lambda: "ok", timeout=5.0)

    check(asyncio.run(fast_run()) == "ok", "未超时应原样返回线程结果")


def real_documents() -> None:
    """samples/ 下的真实文档（跳过 generated/）逐个转换并打印，作为质量基线。"""
    if not os.path.isdir(SAMPLES_DIR):
        return
    real: list[str] = []
    for root, _dirs, files in os.walk(SAMPLES_DIR):
        if os.path.abspath(root).startswith(os.path.abspath(GENERATED_DIR)):
            continue
        for name in files:
            if os.path.splitext(name)[1].lower() in SUPPORTED_EXTENSIONS:
                real.append(os.path.join(root, name))

    print(f"\n===== 真实文档（samples/，共 {len(real)} 个） =====")
    if not real:
        print("  （无：把真实 PDF/Office 放进 samples/ 即可作为质量基线）")
        return
    for path in sorted(real):
        print(f"\n----- {os.path.relpath(path, HERE)} -----")
        show(convert_path(path), text_limit=600)


def main() -> int:
    os.makedirs(GENERATED_DIR, exist_ok=True)

    docx_path = os.path.join(GENERATED_DIR, "chinese_sample.docx")
    xlsx_path = os.path.join(GENERATED_DIR, "chinese_sample.xlsx")
    pptx_path = os.path.join(GENERATED_DIR, "chinese_sample.pptx")
    pdf_path = os.path.join(GENERATED_DIR, "ascii_sample.pdf")
    build_docx(docx_path)
    build_xlsx(xlsx_path)
    build_pptx(pptx_path)
    build_pdf(pdf_path)

    positive_case(
        "docx",
        docx_path,
        ["文档转换自测报告", "全文索引", "向量检索", "中文表格单元"],
        LOC_KIND_NONE,
    )
    positive_case(
        "xlsx",
        xlsx_path,
        ["第一张表", "数据表", "服务器采购", "含三年维保", "活跃用户"],
        LOC_KIND_SHEET,
    )
    positive_case(
        "pptx",
        pptx_path,
        ["季度汇报", "第三季度进展与下季度计划", "检索能力升级"],
        LOC_KIND_SLIDE,
    )
    positive_case(
        "pdf",
        pdf_path,
        ["Alpha page one chonkpilot", "Beta page two chonkpilot"],
        LOC_KIND_PAGE,
    )

    negative_cases()
    lifecycle_cases()
    security_cases()
    timeout_case()
    real_documents()

    print("\n===== 汇总 =====")
    if _failures:
        print(f"  FAILED: {len(_failures)} 项")
        for item in _failures:
            print(f"   - {item}")
        return 1
    print("  ALL PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
